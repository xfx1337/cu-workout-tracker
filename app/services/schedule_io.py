"""Выгрузка и загрузка расписания через Excel.

В файле лежит НЕДЕЛЬНЫЙ шаблон и период его действия. Бот сам разворачивает
шаблон на каждую неделю периода, так что админу не нужно расписывать семестр руками.

Загрузка перезаписывает расписание внутри периода: будущие занятия там удаляются
и создаются заново. Прошедшие не трогаем — на них висит посещаемость.
"""
from __future__ import annotations

import io
from dataclasses import dataclass, field
from datetime import date, datetime, time, timedelta

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import delete, func, select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from app.config import settings
from app.models import Booking, BookingStatus, Training
from app.tz import WEEKDAYS_RU, local_today, now_utc, to_local, to_utc

SHEET = "Расписание"
HEADERS = ["День недели", "Время", "Название", "Тренер", "Описание", "Место", "Минут", "Мест"]
HEADER_ROW = 5          # строки 1-4 заняты заголовком и датами
FIRST_DATA_ROW = HEADER_ROW + 1
MAX_ROWS = 500          # защита от случайного гигантского файла

WEEKDAY_NAMES = {
    "пн": 0, "понедельник": 0,
    "вт": 1, "вторник": 1,
    "ср": 2, "среда": 2,
    "чт": 3, "четверг": 3,
    "пт": 4, "пятница": 4,
    "сб": 5, "суббота": 5,
    "вс": 6, "воскресенье": 6,
}


@dataclass(slots=True)
class Slot:
    weekday: int
    at: time
    title: str
    instructor: str = ""
    description: str = ""
    location: str = ""
    duration_min: int = 60
    capacity: int = 25


@dataclass
class ParsedSchedule:
    starts_on: date | None = None
    ends_on: date | None = None
    slots: list[Slot] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)

    @property
    def ok(self) -> bool:
        return not self.errors and bool(self.slots) and self.starts_on is not None


@dataclass
class ApplyResult:
    created: int = 0
    removed: int = 0
    kept_past: int = 0
    # кому написать, что его запись отменилась: (id в Mattermost, текст про занятие)
    notify: list[tuple[str, str]] = field(default_factory=list)


# ---------------------------------------------------------------- выгрузка

async def export_workbook(session: AsyncSession) -> bytes:
    """Собирает недельный шаблон из будущих занятий."""
    today = local_today()
    stmt = (
        select(Training)
        .where(
            Training.starts_at >= to_utc(datetime.combine(today, time.min)),
            Training.is_cancelled.is_(False),
        )
        .order_by(Training.starts_at)
    )
    items = list(await session.scalars(stmt))

    # схлопываем семестр в одну неделю: одинаковые занятия повторяются каждую неделю
    pattern: dict[tuple[int, str, str], Training] = {}
    for t in items:
        local = to_local(t.starts_at)
        pattern.setdefault((local.weekday(), f"{local:%H:%M}", t.title), t)

    starts_on = to_local(items[0].starts_at).date() if items else today
    ends_on = to_local(items[-1].starts_at).date() if items else today + timedelta(days=120)

    wb = Workbook()
    ws = wb.active
    ws.title = SHEET

    ws["A1"] = "Расписание групповых тренировок"
    ws["A1"].font = Font(bold=True, size=14)

    ws["A2"] = "Начало расписания"
    ws["B2"] = starts_on
    ws["A3"] = "Конец расписания"
    ws["B3"] = ends_on
    for cell in ("A2", "A3"):
        ws[cell].font = Font(bold=True)
    for cell in ("B2", "B3"):
        ws[cell].number_format = "DD.MM.YYYY"

    ws["D2"] = "Ниже — расписание на ОДНУ неделю. Бот повторит его каждую неделю периода."
    ws["D3"] = "Даты выше можно менять. Строки можно добавлять и удалять."
    for cell in ("D2", "D3"):
        ws[cell].font = Font(italic=True, color="666666")

    header_fill = PatternFill("solid", fgColor="222222")
    for col, name in enumerate(HEADERS, start=1):
        cell = ws.cell(row=HEADER_ROW, column=col, value=name)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    row = FIRST_DATA_ROW
    for (weekday, at, _title), t in sorted(pattern.items()):
        ws.cell(row=row, column=1, value=WEEKDAYS_RU[weekday])
        ws.cell(row=row, column=2, value=at)
        ws.cell(row=row, column=3, value=t.title)
        ws.cell(row=row, column=4, value=t.instructor)
        ws.cell(row=row, column=5, value=t.description)
        ws.cell(row=row, column=6, value=t.location)
        ws.cell(row=row, column=7, value=t.duration_min)
        ws.cell(row=row, column=8, value=t.capacity)
        row += 1

    for col, width in zip("ABCDEFGH", (14, 10, 30, 24, 50, 20, 8, 8)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = ws.cell(row=FIRST_DATA_ROW, column=1)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------- разбор

def _as_date(value, label: str, errors: list[str]) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str) and value.strip():
        for fmt in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
            try:
                return datetime.strptime(value.strip(), fmt).date()
            except ValueError:
                continue
    errors.append(f"{label}: не понял дату «{value}». Нужен формат ДД.ММ.ГГГГ.")
    return None


def _as_time(value, row: int, errors: list[str]) -> time | None:
    if isinstance(value, datetime):
        return value.time()
    if isinstance(value, time):
        return value
    if isinstance(value, str) and value.strip():
        raw = value.strip().replace(".", ":")
        for fmt in ("%H:%M", "%H:%M:%S"):
            try:
                return datetime.strptime(raw, fmt).time()
            except ValueError:
                continue
    errors.append(f"Строка {row}: не понял время «{value}». Нужен формат ЧЧ:ММ.")
    return None


def _as_int(value, row: int, label: str, default: int, errors: list[str]) -> int:
    if value is None or (isinstance(value, str) and not value.strip()):
        return default
    try:
        number = int(float(value))
    except (TypeError, ValueError):
        errors.append(f"Строка {row}: «{label}» должно быть числом, а не «{value}».")
        return default
    if number <= 0:
        errors.append(f"Строка {row}: «{label}» должно быть больше нуля.")
        return default
    return number


def _text(value) -> str:
    return str(value).strip() if value is not None else ""


def parse_workbook(data: bytes) -> ParsedSchedule:
    parsed = ParsedSchedule()
    try:
        wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    except Exception as exc:
        parsed.errors.append(f"Не смог открыть файл как Excel: {exc}")
        return parsed

    ws = wb[SHEET] if SHEET in wb.sheetnames else wb.worksheets[0]
    rows = list(ws.iter_rows(min_row=1, max_row=MAX_ROWS, max_col=len(HEADERS), values_only=True))
    wb.close()

    def cell(r: int, c: int):
        if r - 1 < len(rows) and c - 1 < len(rows[r - 1]):
            return rows[r - 1][c - 1]
        return None

    parsed.starts_on = _as_date(cell(2, 2), "Начало расписания", parsed.errors)
    parsed.ends_on = _as_date(cell(3, 2), "Конец расписания", parsed.errors)

    if parsed.starts_on and parsed.ends_on and parsed.ends_on < parsed.starts_on:
        parsed.errors.append("Конец расписания раньше начала.")

    seen: set[tuple[int, time, str]] = set()
    for index in range(FIRST_DATA_ROW, len(rows) + 1):
        raw = rows[index - 1]
        if raw is None or all(v is None or _text(v) == "" for v in raw):
            continue

        row_no = index
        weekday_raw = _text(cell(row_no, 1)).lower()
        weekday = WEEKDAY_NAMES.get(weekday_raw)
        if weekday is None:
            parsed.errors.append(
                f"Строка {row_no}: не понял день недели «{_text(cell(row_no, 1))}». "
                "Пиши Пн, Вт, Ср, Чт, Пт, Сб или Вс."
            )
            continue

        at = _as_time(cell(row_no, 2), row_no, parsed.errors)
        title = _text(cell(row_no, 3))
        if not title:
            parsed.errors.append(f"Строка {row_no}: пустое название тренировки.")
        if at is None or not title:
            continue

        key = (weekday, at, title)
        if key in seen:
            parsed.errors.append(f"Строка {row_no}: такая тренировка уже есть выше.")
            continue
        seen.add(key)

        parsed.slots.append(
            Slot(
                weekday=weekday,
                at=at,
                title=title,
                instructor=_text(cell(row_no, 4)),
                description=_text(cell(row_no, 5)),
                location=_text(cell(row_no, 6)),
                duration_min=_as_int(cell(row_no, 7), row_no, "Минут", 60, parsed.errors),
                capacity=_as_int(cell(row_no, 8), row_no, "Мест", settings.default_capacity, parsed.errors),
            )
        )

    if not parsed.slots and not parsed.errors:
        parsed.errors.append("В файле нет ни одной тренировки.")
    return parsed


# ---------------------------------------------------------------- применение

def expand(parsed: ParsedSchedule) -> list[tuple[Slot, datetime]]:
    """Разворачивает недельный шаблон на весь период. Время возвращаем в UTC."""
    result: list[tuple[Slot, datetime]] = []
    now = now_utc()
    day = parsed.starts_on
    while day <= parsed.ends_on:
        for slot in parsed.slots:
            if slot.weekday != day.weekday():
                continue
            starts_at = to_utc(datetime.combine(day, slot.at))
            if starts_at > now:          # прошедшее не создаём
                result.append((slot, starts_at))
        day += timedelta(days=1)
    return sorted(result, key=lambda pair: pair[1])


async def preview(session: AsyncSession, parsed: ParsedSchedule) -> ApplyResult:
    """Считает, что произойдёт, ничего не меняя."""
    return await _apply(session, parsed, dry_run=True)


async def apply_schedule(session: AsyncSession, parsed: ParsedSchedule) -> ApplyResult:
    return await _apply(session, parsed, dry_run=False)


async def _apply(session: AsyncSession, parsed: ParsedSchedule, *, dry_run: bool) -> ApplyResult:
    result = ApplyResult()
    now = now_utc()

    window_start = to_utc(datetime.combine(parsed.starts_on, time.min))
    window_end = to_utc(datetime.combine(parsed.ends_on + timedelta(days=1), time.min))

    # внутри периода сносим только будущие занятия: на прошедших висит посещаемость
    doomed_stmt = select(Training).where(
        Training.starts_at >= window_start,
        Training.starts_at < window_end,
        Training.starts_at > now,
    )
    doomed = list(await session.scalars(doomed_stmt))
    doomed_ids = [t.id for t in doomed]
    result.removed = len(doomed)
    result.kept_past = await session.scalar(
        select(func.count(Training.id)).where(
            Training.starts_at >= window_start,
            Training.starts_at < window_end,
            Training.starts_at <= now,
        )
    ) or 0

    if doomed_ids:
        bookings = list(
            await session.scalars(
                select(Booking)
                .where(Booking.training_id.in_(doomed_ids), Booking.status == BookingStatus.BOOKED)
                .options(selectinload(Booking.student), selectinload(Booking.training))
            )
        )
        for b in bookings:
            local = to_local(b.training.starts_at)
            result.notify.append(
                (b.student.mm_user_id, f"{b.training.title} — {local:%d.%m} в {local:%H:%M}")
            )

    planned = expand(parsed)
    result.created = len(planned)

    if dry_run:
        return result

    if doomed_ids:
        # чистим записи явно: на SQLite внешние ключи по умолчанию не каскадируют
        await session.execute(delete(Booking).where(Booking.training_id.in_(doomed_ids)))
        await session.execute(delete(Training).where(Training.id.in_(doomed_ids)))

    for slot, starts_at in planned:
        session.add(
            Training(
                title=slot.title,
                description=slot.description,
                instructor=slot.instructor,
                location=slot.location,
                starts_at=starts_at,
                duration_min=slot.duration_min,
                capacity=slot.capacity,
            )
        )
    await session.commit()
    return result

