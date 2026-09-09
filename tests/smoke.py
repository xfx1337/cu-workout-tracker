"""Смоук-тест доменной логики без Telegram.

Запуск:  python -m tests.smoke
Гоняет реальные модели и сервисы на временной SQLite-базе.
"""
from __future__ import annotations

import asyncio
import os
import tempfile
from datetime import timedelta

# По умолчанию гоняем на временной SQLite. Чтобы проверить боевой движок:
#   SMOKE_DATABASE_URL=postgresql+asyncpg://workout:workout@127.0.0.1:5432/workout_test
# Схема в начале прогона сносится, так что база должна быть выделенной под тесты.
DB_PATH = os.path.join(tempfile.mkdtemp(prefix="cu-workout-"), "smoke.db")
os.environ.setdefault("BOT_TOKEN", "test:token")
os.environ["DATABASE_URL"] = os.getenv("SMOKE_DATABASE_URL") or f"sqlite+aiosqlite:///{DB_PATH}"
os.environ["ADMIN_USERNAMES"] = "boss"
os.environ["AUTH_API_URL"] = ""
os.environ["NO_SHOW_LIMIT"] = "3"

from app.auth import resolve_student  # noqa: E402
from app.db import SessionMaker, init_db  # noqa: E402
from app.models import Attendance, Student  # noqa: E402
from app.services import bookings as bookings_svc  # noqa: E402
from app.services import enforcement  # noqa: E402
from app.services import stats as stats_svc  # noqa: E402
from app.services import students as students_svc  # noqa: E402
from app.services import trainings as trainings_svc  # noqa: E402
from app.services.bookings import BookResult, CancelResult  # noqa: E402
from app.tz import now_utc, to_local, week_offset_of  # noqa: E402

PASSED = 0
_DISPATCHER = None


def get_dispatcher():
    """Роутеры — модульные синглтоны, поэтому диспетчер строим один раз на процесс."""
    global _DISPATCHER
    if _DISPATCHER is None:
        from app.bot import build_dispatcher

        _DISPATCHER = build_dispatcher()
    return _DISPATCHER


def check(label: str, condition: bool) -> None:
    global PASSED
    if not condition:
        raise AssertionError(f"FAIL: {label}")
    PASSED += 1
    print(f"  ok  {label}")


class FakeBot:
    """Заглушка вместо aiogram.Bot — собирает отправленные сообщения."""

    def __init__(self) -> None:
        self.sent: list[tuple[int, str]] = []

    async def send_message(self, chat_id: int, text: str, **kwargs) -> None:
        self.sent.append((chat_id, text))


async def make_student(session, tg_id: int) -> Student:
    student = await students_svc.get_or_create(session, tg_id, f"user{tg_id}")
    identity = await resolve_student(tg_id, f"user{tg_id}")
    await students_svc.attach_identity(session, student, identity)
    await students_svc.accept_consent(session, student)
    return student


async def check_race() -> None:
    """Два студента одновременно жмут «записаться» на последнее место."""
    async with SessionMaker() as session:
        training = await trainings_svc.create(
            session, title="Гонка", description="", location="",
            starts_at=now_utc() + timedelta(days=3), duration_min=60, capacity=1,
        )
        a = await make_student(session, 101)
        b = await make_student(session, 102)
        training_id, a_id, b_id = training.id, a.id, b.id

    async def attempt(student_id: int) -> BookResult:
        async with SessionMaker() as s:
            return await bookings_svc.book(
                s, await s.get(Student, student_id), await trainings_svc.by_id(s, training_id)
            )

    results = await asyncio.gather(attempt(a_id), attempt(b_id))
    check("при гонке за последнее место побеждает ровно один",
          sorted(r.value for r in results) == ["full", "ok"])

    async with SessionMaker() as session:
        check("вместимость не превышена", await trainings_svc.taken(session, training_id) == 1)


async def check_schedule_io() -> None:
    """Выгрузка расписания в Excel, правка и загрузка обратно."""
    import io
    from datetime import date, datetime, time

    from openpyxl import load_workbook

    from app.services import schedule_io

    async with SessionMaker() as session:
        # своя изолированная неделя далеко в будущем, чтобы не задеть остальные тесты
        monday = (now_utc() + timedelta(days=120)).date()
        monday -= timedelta(days=monday.weekday())
        base = datetime.combine(monday, datetime.min.time())
        old = await trainings_svc.create(
            session, title="Старое занятие", description="описание", location="Зал",
            instructor="Прежний тренер", starts_at=base.replace(hour=19),
            duration_min=60, capacity=25,
        )
        victim = await make_student(session, 601)
        check("студент записался на занятие, которое потом перезапишут",
              await bookings_svc.book(session, victim, old) is BookResult.OK)

        # --- выгрузка ---
        blob = await schedule_io.export_workbook(session)
        check("выгрузка — это xlsx", blob[:2] == b"PK")
        wb = load_workbook(io.BytesIO(blob))
        ws = wb[schedule_io.SHEET]
        check("в файле есть лист «Расписание»", ws.title == "Расписание")
        check("в B2 дата начала", isinstance(ws["B2"].value, (date, datetime)))
        check("в B3 дата конца", isinstance(ws["B3"].value, (date, datetime)))
        titles = {ws.cell(row=r, column=3).value for r in range(schedule_io.FIRST_DATA_ROW, ws.max_row + 1)}
        check("выгруженный шаблон содержит занятие", "Старое занятие" in titles)
        check("шаблон схлопнут в одну неделю: не больше 7 дней",
              len({ws.cell(row=r, column=1).value for r in range(schedule_io.FIRST_DATA_ROW, ws.max_row + 1)}) <= 7)

        # --- правим файл руками, как это сделает админ ---
        wb2 = load_workbook(io.BytesIO(blob))
        ws2 = wb2[schedule_io.SHEET]
        ws2["B2"] = monday
        ws2["B3"] = monday + timedelta(days=13)          # ровно две недели
        for row in range(ws2.max_row, schedule_io.FIRST_DATA_ROW - 1, -1):
            ws2.delete_rows(row)
        ws2.cell(row=schedule_io.FIRST_DATA_ROW, column=1, value="Ср")
        ws2.cell(row=schedule_io.FIRST_DATA_ROW, column=2, value="18:30")
        ws2.cell(row=schedule_io.FIRST_DATA_ROW, column=3, value="Новое занятие")
        ws2.cell(row=schedule_io.FIRST_DATA_ROW, column=4, value="Новый тренер")
        ws2.cell(row=schedule_io.FIRST_DATA_ROW, column=5, value="Новое описание")
        ws2.cell(row=schedule_io.FIRST_DATA_ROW, column=6, value="Зал 2")
        ws2.cell(row=schedule_io.FIRST_DATA_ROW, column=7, value=90)
        ws2.cell(row=schedule_io.FIRST_DATA_ROW, column=8, value=12)
        edited = io.BytesIO()
        wb2.save(edited)

        parsed = schedule_io.parse_workbook(edited.getvalue())
        check("файл разобрался без ошибок", parsed.ok and not parsed.errors)
        check("период прочитан", (parsed.starts_on, parsed.ends_on) == (monday, monday + timedelta(days=13)))
        check("в шаблоне одно занятие", len(parsed.slots) == 1)
        slot = parsed.slots[0]
        check("день недели разобран", slot.weekday == 2)
        check("время разобрано", slot.at == time(18, 30))
        check("длительность и мест разобраны", (slot.duration_min, slot.capacity) == (90, 12))

        # --- предпросмотр ничего не меняет ---
        before = len(await trainings_svc.for_week(session, monday))
        plan = await schedule_io.preview(session, parsed)
        check("предпросмотр обещает 2 занятия на 2 недели", plan.created == 2)
        check("предпросмотр видит, что старое удалится", plan.removed >= 1)
        check("предпросмотр предупреждает про слетающие записи", len(plan.notify) == 1)
        check("предпросмотр ничего не изменил",
              len(await trainings_svc.for_week(session, monday)) == before)

        # --- применяем ---
        result = await schedule_io.apply_schedule(session, parsed)
        check("создано ровно по одному занятию на неделю", result.created == 2)

        week1 = await trainings_svc.for_week(session, monday)
        week2 = await trainings_svc.for_week(session, monday + timedelta(days=7))
        check("на первой неделе новое занятие", [t.title for t in week1] == ["Новое занятие"])
        check("на второй неделе оно тоже есть", [t.title for t in week2] == ["Новое занятие"])
        check("старое занятие удалено", "Старое занятие" not in {t.title for t in week1})
        check("детали занятия доехали",
              (week1[0].instructor, week1[0].location, week1[0].capacity)
              == ("Новый тренер", "Зал 2", 12))
        check("занятие встало в среду", to_local(week1[0].starts_at).weekday() == 2)
        check("за пределы периода не вылезли",
              not await trainings_svc.for_week(session, monday + timedelta(days=14)))
        check("запись студента исчезла вместе с занятием",
              not await bookings_svc.active_for_student(session, victim.id))

        # --- повторная загрузка того же файла не плодит дубли ---
        again = await schedule_io.apply_schedule(session, parsed)
        check("повторная загрузка не дублирует", again.created == 2)
        check("на неделе по-прежнему одно занятие",
              len(await trainings_svc.for_week(session, monday)) == 1)


async def check_schedule_io_errors() -> None:
    """Кривой файл должен внятно ругаться, а не падать."""
    import io

    from openpyxl import Workbook

    from app.services import schedule_io

    def build(rows, start="01.09.2026", end="30.09.2026") -> bytes:
        wb = Workbook()
        ws = wb.active
        ws.title = schedule_io.SHEET
        ws["B2"] = start
        ws["B3"] = end
        for i, row in enumerate(rows):
            for col, value in enumerate(row, start=1):
                ws.cell(row=schedule_io.FIRST_DATA_ROW + i, column=col, value=value)
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    parsed = schedule_io.parse_workbook(b"not an excel file at all")
    check("мусор вместо файла не роняет бота", not parsed.ok and parsed.errors)

    parsed = schedule_io.parse_workbook(build([("Понедельник", "8:30", "Утро")]))
    check("полное название дня недели тоже понимается", parsed.ok and parsed.slots[0].weekday == 0)

    parsed = schedule_io.parse_workbook(build([("Хрень", "18:30", "Занятие")]))
    check("непонятный день недели — ошибка с номером строки",
          not parsed.ok and "Строка 6" in parsed.errors[0])

    parsed = schedule_io.parse_workbook(build([("Пн", "четверть восьмого", "Занятие")]))
    check("непонятное время — ошибка", not parsed.ok and "время" in parsed.errors[0])

    parsed = schedule_io.parse_workbook(build([("Пн", "18:30", "")]))
    check("пустое название — ошибка", not parsed.ok and "название" in parsed.errors[0])

    parsed = schedule_io.parse_workbook(build([("Пн", "18:30", "А"), ("Пн", "18:30", "А")]))
    check("дубль в шаблоне — ошибка", not parsed.ok and "уже есть" in parsed.errors[0])

    parsed = schedule_io.parse_workbook(build([("Пн", "18:30", "А")], start="30.09.2026", end="01.09.2026"))
    check("конец раньше начала — ошибка", not parsed.ok)

    parsed = schedule_io.parse_workbook(build([]))
    check("пустой шаблон — ошибка", not parsed.ok and "ни одной тренировки" in parsed.errors[0])

    parsed = schedule_io.parse_workbook(build([("Пн", "18:30", "А", "", "", "", "ноль", -5)]))
    check("нечисловая длительность и отрицательные места — ошибки", len(parsed.errors) == 2)


async def check_cancelled_hidden() -> None:
    """Отменённое занятие остаётся на картинке, но исчезает из всего остального."""
    from datetime import datetime

    from app.handlers.student import _day_view, _training_caption, _week_view
    from app.keyboards import training_kb
    from app.render import status_of

    async with SessionMaker() as session:
        student = await make_student(session, 501)

        monday = (now_utc() + timedelta(days=14)).date()
        monday -= timedelta(days=monday.weekday())
        base = datetime.combine(monday, datetime.min.time())
        alive = await trainings_svc.create(
            session, title="Живое занятие", description="", location="", instructor="Тренер",
            starts_at=base.replace(hour=10), duration_min=60, capacity=25,
        )
        doomed = await trainings_svc.create(
            session, title="Отменяемое занятие", description="", location="", instructor="Тренер",
            starts_at=base.replace(hour=12), duration_min=60, capacity=25,
        )
        lonely = await trainings_svc.create(   # единственное занятие вторника
            session, title="Одинокое занятие", description="", location="", instructor="Тренер",
            starts_at=(base + timedelta(days=1)).replace(hour=10), duration_min=60, capacity=25,
        )
        offset = week_offset_of(alive.starts_at)

        check("до отмены студент видит оба занятия дня",
              len((await _day_view(session, student, offset, 0))[1].inline_keyboard) == 3)

        await trainings_svc.cancel(session, doomed)
        await trainings_svc.cancel(session, lonely)

        # картинка по-прежнему рисует отменённое — статусом
        png, _, week_kb_after = await _week_view(session, student, offset)
        check("отменённое занятие остаётся на картинке", png[:8] == b"\x89PNG\r\n\x1a\n")
        check("на картинке у него статус «отменена»",
              status_of(doomed, 0, set(), now_utc()) == "cancelled")

        # а из кнопок и списков пропадает
        day_text, day_kb_after = await _day_view(session, student, offset, 0)
        check("в списке дня отменённого занятия нет", "Отменяемое занятие" not in day_text)
        check("живое занятие в списке дня осталось", "Живое занятие" in day_text)
        check("кнопка отменённого занятия убрана", len(day_kb_after.inline_keyboard) == 2)

        day_labels = [b.text for row in week_kb_after.inline_keyboard for b in row]
        check("день, где всё отменено, не предлагается", "Вт" not in day_labels)
        check("день с живым занятием предлагается", "Пн" in day_labels)

        # карточка: записаться нельзя
        text, is_booked = await _training_caption(session, student, doomed)
        check("в карточке видно, что занятие отменено", "ОТМЕНЕНА" in text)
        check("в карточке сказано, что записаться нельзя", "записаться нельзя" in text)
        check("студент не считается записанным на отменённое", is_booked is False)

        buttons = [b.text for row in training_kb(doomed, False, offset, 0).inline_keyboard for b in row]
        check("кнопки «Записаться» на отменённом занятии нет",
              not any("Записаться" in b for b in buttons))
        check("кнопка «назад» при этом осталась", any("⬅️" in b for b in buttons))
        alive_buttons = [b.text for row in training_kb(alive, False, offset, 0).inline_keyboard for b in row]
        check("у живого занятия кнопка записи на месте",
              any("Записаться" in b for b in alive_buttons))

        # и сервер тоже отказывает, даже если нажать старую кнопку
        check("запись на отменённое занятие отбивается",
              await bookings_svc.book(session, student, doomed) is BookResult.CANCELLED)


async def check_render() -> None:
    """Картинка расписания: строится, читается, статусы считаются верно."""
    import io
    from datetime import date, datetime, timedelta

    from PIL import Image

    from app.models import Training
    from app.render import plural_ru, render_week, status_of, week_title
    from app.services import trainings as tsvc

    check("склонение «мест»", [plural_ru(n, "место", "места", "мест") for n in (1, 3, 5, 11, 21, 22)]
          == ["место", "места", "мест", "мест", "место", "места"])
    check("заголовок недели в одном месяце", week_title(date(2026, 9, 14), date(2026, 9, 20)) == "14 – 20 сентября")
    check("заголовок недели на стыке месяцев",
          week_title(date(2026, 9, 28), date(2026, 10, 4)) == "28 сентября – 4 октября")

    future = now_utc() + timedelta(days=1)
    t = Training(id=1, title="Йога", description="", instructor="Кто-то", location="",
                 starts_at=future, duration_min=60, capacity=25, is_cancelled=False)
    check("свободная тренировка", status_of(t, 3, set(), now_utc()) == "free")
    check("своя запись важнее свободных мест", status_of(t, 3, {1}, now_utc()) == "booked")
    check("забитая тренировка", status_of(t, 25, set(), now_utc()) == "full")
    t.is_cancelled = True
    check("отменённая тренировка", status_of(t, 3, {1}, now_utc()) == "cancelled")
    t.is_cancelled = False
    t.starts_at = now_utc() - timedelta(hours=3)
    check("прошедшая тренировка", status_of(t, 3, set(), now_utc()) == "past")

    # реальная недельная выборка из базы
    async with SessionMaker() as session:
        monday = (now_utc() + timedelta(days=7)).date()
        monday -= timedelta(days=monday.weekday())
        for i, hour in enumerate((9, 18)):
            await tsvc.create(
                session, title=f"Занятие {i}", description="описание", location="Зал",
                instructor="Тренер Тренерович",
                starts_at=datetime.combine(monday, datetime.min.time()).replace(hour=hour),
                duration_min=60, capacity=25,
            )
        week_items = await tsvc.for_week(session, monday)
        check("выборка по неделе находит занятия", len(week_items) >= 2)
        day_items = await tsvc.for_day(session, monday)
        check("выборка по дню находит занятия", len(day_items) >= 2)
        check("занятие следующей недели не попадает в текущую",
              all(monday <= to_local(x.starts_at).date() < monday + timedelta(days=7) for x in week_items))

        png = render_week(week_items, {week_items[0].id: 25}, {week_items[1].id}, monday)

    check("на выходе PNG", png[:8] == b"\x89PNG\r\n\x1a\n")
    img = Image.open(io.BytesIO(png))
    check("картинка нужной ширины", img.width == 1240)
    check("картинка не выродилась по высоте", 400 < img.height < 4000)
    check("картинка не пустая", len(img.getcolors(maxcolors=100000) or []) > 5)

    check("пустая неделя тоже рисуется без падения", render_week([], {}, set(), monday)[:8] == b"\x89PNG\r\n\x1a\n")


class RecordingSession:
    """Подменяет сетевую сессию aiogram: ничего не шлёт, только записывает вызовы."""

    def __init__(self) -> None:
        self.requests: list = []

    async def close(self) -> None:
        pass

    async def stream_content(self, *args, **kwargs):  # pragma: no cover
        yield b""

    async def make_request(self, bot, method, timeout=None):
        from datetime import datetime

        from aiogram.types import Chat, Message, User

        self.requests.append(method)
        name = type(method).__name__
        if name.startswith(("Send", "Edit")):
            return Message(
                message_id=999, date=datetime.now(),
                chat=Chat(id=1, type="private"),
                from_user=User(id=1, is_bot=True, first_name="bot"),
                text="stub",
            )
        if name == "GetMe":
            return User(id=1, is_bot=True, first_name="bot", username="bot")
        return True


async def check_dispatch() -> None:
    """Гоняем настоящие апдейты через диспетчер.

    Ловит то, что не видно при простой сборке: например, inner-мидлварь
    выполняется ПОСЛЕ фильтров, и фильтр IsAdmin не получает admin.
    """
    from datetime import datetime

    from aiogram import Bot
    from aiogram.client.default import DefaultBotProperties
    from aiogram.client.session.base import BaseSession
    from aiogram.dispatcher.event.bases import UNHANDLED
    from aiogram.enums import ParseMode
    from aiogram.types import CallbackQuery, Chat, Message, Update, User

    from app.keyboards import NavCB

    class Session(RecordingSession, BaseSession):
        def __init__(self) -> None:
            BaseSession.__init__(self)
            RecordingSession.__init__(self)

    session = Session()
    bot = Bot(
        token="123456:AAFakeTokenForTests",
        session=session,
        default=DefaultBotProperties(parse_mode=ParseMode.HTML),
    )
    dp = get_dispatcher()

    admin_user = User(id=999, is_bot=False, first_name="Boss", username="boss")
    plain_user = User(id=1, is_bot=False, first_name="Stu", username="user1")
    counter = [0]

    def message_update(user: User, text: str) -> Update:
        counter[0] += 1
        return Update(
            update_id=counter[0],
            message=Message(
                message_id=counter[0], date=datetime.now(),
                chat=Chat(id=user.id, type="private"), from_user=user, text=text,
            ),
        )

    def callback_update(user: User, data: str) -> Update:
        counter[0] += 1
        return Update(
            update_id=counter[0],
            callback_query=CallbackQuery(
                id=str(counter[0]), from_user=user, chat_instance="ci", data=data,
                message=Message(
                    message_id=counter[0], date=datetime.now(),
                    chat=Chat(id=user.id, type="private"),
                    from_user=User(id=1, is_bot=True, first_name="bot"), text="экран",
                ),
            ),
        )

    async def feed(update: Update) -> tuple[object, list[str]]:
        session.requests.clear()
        result = await dp.feed_update(bot, update)
        return result, [type(r).__name__ for r in session.requests]

    result, calls = await feed(message_update(admin_user, "/admin"))
    check("админ открывает админку командой", result is not UNHANDLED and "SendMessage" in calls)

    result, calls = await feed(
        callback_update(admin_user, NavCB(action="admin", offset=0).pack())
    )
    check("админка открывается кнопкой из расписания",
          result is not UNHANDLED and "SendMessage" in calls)

    result, _ = await feed(message_update(plain_user, "/admin"))
    check("обычного студента в админку не пускают", result is UNHANDLED)

    result, calls = await feed(message_update(plain_user, "/start"))
    check("/start авторизованному студенту сразу шлёт расписание картинкой",
          result is not UNHANDLED and "SendPhoto" in calls)

    result, calls = await feed(message_update(plain_user, "/my"))
    check("«Мои записи» открываются командой", result is not UNHANDLED and "SendMessage" in calls)

    # раздел «Тренировки» в админке — та же картинка недели, а не простыня кнопок
    from app.keyboards import AdmCB, AdmWeekCB

    result, calls = await feed(callback_update(admin_user, AdmCB(action="trainings").pack()))
    check("«Тренировки» в админке отдают картинку недели",
          result is not UNHANDLED and "SendPhoto" in calls)

    result, calls = await feed(callback_update(admin_user, AdmWeekCB(action="week", offset=1).pack()))
    check("админ листает недели", result is not UNHANDLED and "SendPhoto" in calls)

    result, calls = await feed(callback_update(admin_user, AdmCB(action="export_schedule").pack()))
    check("выгрузка расписания отдаёт файл",
          result is not UNHANDLED and "SendDocument" in calls)

    result, calls = await feed(callback_update(admin_user, AdmCB(action="import_schedule").pack()))
    check("загрузка расписания просит файл", result is not UNHANDLED)
    result, calls = await feed(message_update(admin_user, "не файл, а текст"))
    check("вместо файла прислали текст — бот объясняет, что ждёт xlsx",
          result is not UNHANDLED and "SendMessage" in calls)
    await feed(callback_update(admin_user, AdmCB(action="menu").pack()))   # выходим из состояния

    result, calls = await feed(
        callback_update(admin_user, AdmWeekCB(action="day", offset=0, weekday=0).pack())
    )
    check("админ открывает день", result is not UNHANDLED)

    # выход из админки: текстовый экран удаляется, взамен приходит картинка
    result, calls = await feed(
        callback_update(admin_user, NavCB(action="schedule", offset=0).pack())
    )
    check("из админки выходим к расписанию", result is not UNHANDLED and "SendPhoto" in calls)
    check("текстовый экран админки при этом убирается", "DeleteMessage" in calls)

    await bot.session.close()


async def check_wiring() -> None:
    """Проверяем, что бот вообще собирается: роутеры, фильтры, клавиатуры."""
    from app.keyboards import (
        AdmCB, AdmTrainingCB, AttendCB, PollCB, TrainingCB, admin_menu_kb, attendance_kb,
        schedule_kb, student_card_kb, training_kb,
    )
    from app.models import Booking, Training

    dp = get_dispatcher()
    check("диспетчер собирается", dp is not None)
    check("зарегистрированы 3 роутера", len(dp.sub_routers) == 3)
    check("апдейты резолвятся", "message" in dp.resolve_used_update_types())

    # callback data должна паковаться и распаковываться без потерь
    packed = TrainingCB(action="book", training_id=42).pack()
    check("callback data round-trip", TrainingCB.unpack(packed).training_id == 42)
    check("длина callback data в лимите телеграма (64 байта)",
          max(len(x) for x in (
              packed,
              AdmCB(action="del_admin", obj_id=999999).pack(),
              AdmTrainingCB(action="cancel_confirm", training_id=999999).pack(),
              AttendCB(booking_id=999999, value="toggle").pack(),
              PollCB(booking_id=999999, came=False).pack(),
          )) <= 64)

    # клавиатуры собираются на реальных объектах
    training = Training(id=1, title="Тест", description="", location="", starts_at=now_utc(),
                        duration_min=60, capacity=25, is_cancelled=False)
    student = Student(id=1, tg_user_id=1, tg_username="u", no_show_count=1, is_banned=False)
    booking = Booking(id=1, training_id=1, student_id=1, attendance=Attendance.UNKNOWN)
    booking.student = student

    check("расписание рисуется", schedule_kb([(training, 3)], set()).inline_keyboard[0][0].text.startswith("🟢"))
    check("занятая тренировка помечена красным",
          schedule_kb([(training, 25)], set()).inline_keyboard[0][0].text.startswith("🔴"))
    check("своя запись помечена галочкой",
          schedule_kb([(training, 3)], {1}).inline_keyboard[0][0].text.startswith("✅"))
    check("карточка тренировки", len(training_kb(training, False).inline_keyboard) == 2)
    check("экран посещаемости", len(attendance_kb([booking], 1).inline_keyboard) == 2)
    check("карточка студента даёт сброс пропусков", len(student_card_kb(student).inline_keyboard) == 3)
    adm_menu = admin_menu_kb()
    adm_labels = [b.text for row in adm_menu.inline_keyboard for b in row]
    check("меню админки", len(adm_menu.inline_keyboard) == 6)
    check("из админки есть выход к расписанию", "расписанию" in adm_menu.inline_keyboard[-1][0].text)
    check("в меню есть выгрузка расписания", any("Выгрузить расписание" in b for b in adm_labels))
    check("в меню есть загрузка расписания", any("Загрузить расписание" in b for b in adm_labels))
    check("поштучное добавление тренировки осталось", any("Новая тренировка" in b for b in adm_labels))

    from app.keyboards import (
        AdmWeekCB, NavCB, WeekCB, admin_day_kb, admin_week_kb, day_kb, week_kb,
    )
    wk = week_kb(0, [(0, False), (1, True)], has_next=True)
    check("в сетке недели день со своей записью помечен", wk.inline_keyboard[0][1].text == "✅ Вт")
    check("у дня недели нет лишних цифр", wk.inline_keyboard[0][0].text == "Пн")
    def week_nav(markup) -> list[str]:
        return [b.text for row in markup.inline_keyboard for b in row if "неделя" in b.text]

    check("на первой неделе нет кнопки «назад»", week_nav(wk) == ["След. неделя →"])
    check("на дальней неделе есть обе кнопки навигации",
          week_nav(week_kb(2, [(0, False)], has_next=True)) == ["← Пред. неделя", "След. неделя →"])
    check("на последней неделе кнопки «вперёд» нет",
          week_nav(week_kb(2, [(0, False)], has_next=False)) == ["← Пред. неделя"])
    check("callback дня round-trip", WeekCB.unpack(
        WeekCB(action="day", offset=3, weekday=4).pack()).weekday == 4)
    dk = day_kb(1, 2, [(training, 25, False)])
    check("в списке дня забитая тренировка помечена красным", dk.inline_keyboard[0][0].text.startswith("🔴"))
    check("callback тренировки помнит день и неделю",
          TrainingCB.unpack(dk.inline_keyboard[0][0].callback_data).weekday == 2)

    # в админке выбор тренировки такой же: неделя -> день -> занятие
    awk = admin_week_kb(0, [0, 2, 4], has_prev=False, has_next=True)
    check("админская сетка даёт кнопки дней", [b.text for b in awk.inline_keyboard[0]] == ["Пн", "Ср", "Пт"])
    check("админ листает недели вперёд",
          [b.text for b in awk.inline_keyboard[1]] == ["След. неделя →"])
    check("админ уходит в прошлое за посещаемостью",
          len(admin_week_kb(-1, [0], has_prev=True, has_next=True).inline_keyboard[1]) == 2)
    check("из админской сетки есть выход в меню",
          awk.inline_keyboard[-1][0].text.endswith("В админку"))
    check("callback админского дня round-trip",
          AdmWeekCB.unpack(AdmWeekCB(action="day", offset=-3, weekday=5).pack()).offset == -3)

    adk = admin_day_kb(2, 3, [(training, 25)])
    check("в админском списке дня видно занятость", "25/25" in adk.inline_keyboard[0][0].text)
    check("админский callback тренировки помнит неделю и день",
          (lambda c: (c.offset, c.weekday))(AdmTrainingCB.unpack(adk.inline_keyboard[0][0].callback_data)) == (2, 3))

    # интерфейс полностью инлайновый: reply-клавиатур в коде не осталось
    import app.keyboards as kbmod
    check("reply-клавиатура нигде не собирается",
          not any("ReplyKeyboard" in name for name in dir(kbmod)))

    all_buttons = [b for row in wk.inline_keyboard for b in row]
    check("под расписанием есть «Мои записи»", any("Мои записи" in b.text for b in all_buttons))
    check("обычному студенту админку не показываем",
          not any("Админка" in b.text for b in all_buttons))
    admin_wk = week_kb(0, [(0, False)], has_next=False, is_admin=True)
    admin_buttons = [b for row in admin_wk.inline_keyboard for b in row]
    check("админу кнопка админки видна", any("Админка" in b.text for b in admin_buttons))
    check("кнопка админки ведёт на NavCB",
          NavCB.unpack(next(b for b in admin_buttons if "Админка" in b.text).callback_data).action == "admin")


async def main() -> None:
    from app.db import engine
    from app.models import Base

    print(f"движок: {engine.url.get_backend_name()}\n")
    async with engine.begin() as conn:      # прогон всегда с чистой схемы
        await conn.run_sync(Base.metadata.drop_all)
    await init_db()
    bot = FakeBot()

    async with SessionMaker() as session:
        # --- админы из .env ---
        await students_svc.bootstrap_admins(session)
        admin = await students_svc.is_admin(session, 999, "BOSS")
        check("суперадмин подхватывается из ADMIN_USERNAMES по username", admin is not None)
        check("tg_user_id админа запоминается при первом заходе", admin.tg_user_id == 999)

        # --- авторизация-заглушка ---
        student = await make_student(session, 1)
        check("заглушка авторизации выдаёт external_student_id", student.external_student_id == "stub-1")
        check("согласие со справкой/ТБ сохранено", student.consent_accepted_at is not None)

        # --- тренировка с лимитом мест ---
        training = await trainings_svc.create(
            session,
            title="Функциональная",
            description="Круговая",
            location="Зал",
            starts_at=now_utc() + timedelta(days=1),
            duration_min=60,
            capacity=2,
        )

        check("первая запись проходит", await bookings_svc.book(session, student, training) is BookResult.OK)
        check("повторная запись отбивается", await bookings_svc.book(session, student, training) is BookResult.ALREADY)

        second = await make_student(session, 2)
        check("вторая запись занимает последнее место",
              await bookings_svc.book(session, second, training) is BookResult.OK)

        third = await make_student(session, 3)
        check("третьего не пускаем — мест нет",
              await bookings_svc.book(session, third, training) is BookResult.FULL)
        check("занято ровно capacity", await trainings_svc.taken(session, training.id) == 2)

        # --- отмена освобождает место ---
        check("отмена проходит", await bookings_svc.cancel(session, second, training) is CancelResult.OK)
        check("место освободилось", await trainings_svc.taken(session, training.id) == 1)
        check("после отмены третий записывается",
              await bookings_svc.book(session, third, training) is BookResult.OK)

        cancelled = await bookings_svc.get_booking(session, second.id, training.id)
        check("отменённая запись помечена как уважительная", cancelled.attendance is Attendance.EXCUSED)
        check("отмена не даёт страйк", second.no_show_count == 0)

        check("освободившееся место занято — второму уже нельзя вернуться",
              await bookings_svc.book(session, second, training) is BookResult.FULL)

        # --- участники и отмена тренировки админом ---
        people = await trainings_svc.participants(session, training.id)
        check("в участниках двое", len(people) == 2)

        # --- страйки за пропуски ---
        loafer = await make_student(session, 10)
        for i in range(3):
            past = await trainings_svc.create(
                session,
                title=f"Прошедшая {i}",
                description="",
                location="",
                starts_at=now_utc() - timedelta(days=i + 1),
                duration_min=60,
                capacity=25,
            )
            past.starts_at = now_utc() + timedelta(days=1)  # чтобы book() пропустил
            await session.commit()
            await bookings_svc.book(session, loafer, past)
            booking = await bookings_svc.get_booking(session, loafer.id, past.id)
            await session.refresh(booking, ["student"])
            await enforcement.mark_no_show(session, bot, booking, count_strike=True)

        check("после 3 пропусков страйков ровно 3", loafer.no_show_count == 3)
        check("после 3 пропусков автоблокировка", loafer.is_banned is True)
        check("причина блокировки — автоматическая", loafer.ban_reason == "3 пропуска без отмены записи")
        check("студент получил предупреждения и уведомление о блокировке", len(bot.sent) == 3)

        future = await trainings_svc.create(
            session, title="Ещё одна", description="", location="",
            starts_at=now_utc() + timedelta(days=2), duration_min=60, capacity=25,
        )
        check("заблокированный не может записаться",
              await bookings_svc.book(session, loafer, future) is BookResult.BANNED)

        # --- отметка «был» снимает страйк ---
        last_booking = await bookings_svc.get_booking(session, loafer.id, past.id)
        await enforcement.mark_attended(session, last_booking)
        await session.refresh(loafer)
        check("отметка «был» уменьшает счётчик пропусков", loafer.no_show_count == 2)
        check("и снимает автоблокировку", loafer.is_banned is False)

        # --- ручной бан админом ---
        await students_svc.set_ban(session, third, True, "решение администратора")
        check("ручной бан работает", third.is_banned is True)
        banned = await students_svc.banned_list(session)
        check("забаненный попадает в список", third.id in {s.id for s in banned})
        await students_svc.set_ban(session, third, False)
        check("разбан работает", third.is_banned is False)

        # --- поиск ---
        found = await students_svc.search(session, "user1")
        check("поиск студента по username находит", any(s.tg_user_id == 1 for s in found))

        # --- отмена тренировки админом ---
        affected = await trainings_svc.cancel(session, training)
        check("отмена тренировки возвращает список участников", len(affected) == 2)
        check("тренировка помечена отменённой", training.is_cancelled is True)
        check("после отмены записаться нельзя",
              await bookings_svc.book(session, student, training) is BookResult.CANCELLED)

        # --- статистика и выгрузка ---
        overview = await stats_svc.overview(session)
        check("статистика считает студентов", overview["total_students"] >= 4)
        csv_bytes = await stats_svc.export_csv(session)
        check("CSV выгружается и не пустой", len(csv_bytes) > 100)
        check("CSV с заголовком", b"booking_id" in csv_bytes)

    print()
    await check_race()
    await check_cancelled_hidden()
    await check_schedule_io()
    await check_schedule_io_errors()
    await check_render()
    await check_wiring()
    await check_dispatch()
    print(f"\nВсе проверки пройдены: {PASSED}")


if __name__ == "__main__":
    asyncio.run(main())
